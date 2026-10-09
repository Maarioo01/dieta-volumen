#!/usr/bin/env python3
"""Genera data.js (los datos de la app) a partir del Excel DIETA-VOLUMEN.xlsx.

Uso:
    python3 scripts/excel_a_datos.py DIETA-VOLUMEN.xlsx > data.js

Lee las pestañas Macros-<Persona>, DESAYUNOS, ALMUERZOS, COMIDAS Y CENAS,
MERIENDAS y MENU. La pestaña "Registro Nutricional" no se usa.

- Los totales de comidas y cenas NO se copian del Excel: la app los calcula
  sumando los ingredientes (así siempre cuadran con las cantidades).
- Desayunos/almuerzos/meriendas en formato "fila de ingredientes + Kcal/
  Proteínas/Hidratos/Grasas" usan el total que pone el Excel.
- Los nombres bonitos de cada plato y su relación con el MENU están en
  NOMBRES / ALIAS (abajo). Un plato nuevo sin entrada aquí sale con el
  nombre del bloque del Excel en formato "Tipo frase".
"""
import datetime
import json
import re
import sys
import unicodedata

import openpyxl

# ---------------------------------------------------------------- config
# (tipo de comida, título del bloque normalizado) -> nombre en la app
NOMBRES = {
    ("desayuno", "opcion 1"): "Tortilla con tostadas y jamón",
    ("almuerzo", "opcion 1"): "Proteína con fruta",
    ("comida", "macarrones"): "Macarrones",
    ("comida", "ensal. garban."): "Ensalada de garbanzos",
    ("comida", "lomo + calab."): "Lomo con calabacín",
    ("comida", "arroz y lentejas"): "Arroz y lentejas",
    ("comida", "salmon"): "Salmón",
    ("comida", "spaguet. + pavo"): "Espaguetis con pavo",
    ("comida", "espinacas"): "Garbanzos con espinacas y pavo",
    ("cena", "hamburguesa"): "Hamburguesa",
    ("cena", "alitas"): "Alitas de pollo",
    ("cena", "ensal. garban."): "Ensalada de garbanzos",
    ("cena", "ensalada pasta"): "Ensalada de pasta",
    ("cena", "pechuga rellen."): "Pechuga rellena",
    ("cena", "spaguet. + pavo"): "Espaguetis con pavo",
    ("cena", "espinacas"): "Garbanzos con espinacas y pavo",
}
# Variantes por persona dentro de una misma opción: "Opción 1 (Plátano)"
VARIANTE_NOMBRE = {
    ("almuerzo", "opcion 1"): "Proteína con {v}",
}
# Textos del MENU que no coinciden con el nombre en la app
ALIAS = {
    "spaguettis con pavo": "Espaguetis con pavo",
    "proteina y mandarinas / platano": "Proteína con fruta",
}
TIPOS = ["desayuno", "almuerzo", "comida", "merienda", "cena"]


# ---------------------------------------------------------------- helpers
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", norm(s)).strip("-")


def clean(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()


def num(v):
    """Números del Excel: 13.2, '13,2g', '0.5', y celdas que Excel convirtió en fecha/hora."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    if isinstance(v, datetime.datetime):  # p. ej. 21,5 -> 21/01/1900 12:00
        d = v - datetime.datetime(1899, 12, 31)
        return round(d.days + d.seconds / 86400, 4)
    if isinstance(v, datetime.time):  # p. ej. 0,2 -> 04:48
        return round((v.hour * 3600 + v.minute * 60 + v.second) / 86400, 4)
    s = clean(v).lower().replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    return round(float(s), 4) if s else None


def cantidad(v):
    """'160 g', '13g (nuevo)' -> ('13 g', 'nuevo')."""
    s = clean(v)
    nota = None
    m = re.search(r"\s*\(([^)0-9]*)\)\s*$", s)
    if m:
        nota = m.group(1).strip() or None
        s = s[: m.start()].strip()
    s = re.sub(r"(\d)\s*(g|ml|kg)\b", r"\1 \2", s)
    return s, nota


def nombre_receta(tipo, titulo):
    t = norm(titulo)
    if (tipo, t) in NOMBRES:
        return NOMBRES[(tipo, t)]
    return clean(titulo).capitalize()


def merged_value(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def personas_por_columna(ws, header_row=2):
    """[(col, 'mario'), ...] a partir de la fila con MARIO / ALBA."""
    out = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and v.strip():
            out.append((c, slug(v)))
    return out


def persona_en(cols, c):
    p = None
    for cc, pid in cols:
        if cc <= c:
            p = pid
    return p


# ---------------------------------------------------------------- lectores
def leer_objetivos(wb):
    objetivos, personas = {}, []
    for ws in wb.worksheets:
        if not norm(ws.title).startswith("macros-"):
            continue
        pid = slug(ws.title.split("-", 1)[1])
        personas.append({"id": pid, "nombre": clean(ws.title.split("-", 1)[1])})
        tabla = {}
        for row in ws.iter_rows():
            etiqueta = row[0].value
            if not isinstance(etiqueta, str):
                continue
            e = norm(etiqueta)
            if e.startswith("desayuno"):
                tipo = "desayuno"
            elif e.startswith("post-gym") or "media manana" in e:
                tipo = "almuerzo"
            elif e.startswith("merienda"):
                tipo = "merienda"
            elif e.startswith("cena"):
                tipo = "cena"
            elif e.startswith("total"):
                tipo = "total"
            elif e == "comida" and isinstance(row[1].value, str) and "prot" in norm(row[1].value):
                continue  # cabecera
            elif e.startswith("comida"):
                tipo = "comida"
            else:
                continue
            vals = [row[i].value for i in range(1, 5)]
            tabla[tipo] = {
                "etiqueta": clean(etiqueta),
                "p": rango(vals[0]),
                "hc": rango(vals[1]),
                "g": rango(vals[2]),
                "kcalPct": rango(vals[3], pct=True),
            }
        objetivos[pid] = tabla
    return personas, objetivos


def rango(v, pct=False):
    """'25–30' -> {min,max,texto}; '≤10' -> {max}; '≈130 g' -> {aprox}."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        n = v * 100 if pct and v <= 1 else v
        return {"aprox": round(n, 1), "texto": f"{round(n):g}"}
    s = clean(v).replace("%", "").replace("g", "").strip()
    texto = s
    nums = [float(x.replace(",", ".")) for x in re.findall(r"\d+(?:[.,]\d+)?", s)]
    if "≤" in s or "<" in s:
        return {"max": nums[0], "texto": texto}
    if "≈" in s or "~" in s:
        return {"aprox": nums[0], "texto": texto}
    if len(nums) >= 2:
        return {"min": nums[0], "max": nums[1], "texto": texto}
    if nums:
        return {"aprox": nums[0], "texto": texto}
    return None


def leer_opciones_en_fila(ws, tipo, recetas):
    """DESAYUNOS / ALMUERZOS: 'Opción N' + fila de ingredientes + Kcal/Prot/Hidr/Grasas."""
    cols = personas_por_columna(ws)
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and norm(c.value).startswith("opcion")):
                continue
            pid = persona_en(cols, c.column)
            titulo = clean(c.value)
            m = re.match(r"(.*?)\s*\((.+)\)\s*$", titulo)
            base, variante = (m.group(1), m.group(2)) if m else (titulo, None)
            hr, vr = c.row + 1, c.row + 2
            ingredientes, total = [], None
            col = c.column
            while col <= ws.max_column:
                h = ws.cell(hr, col).value
                if h is None:
                    break
                if norm(h) == "kcal":
                    total = {
                        "kcal": num(ws.cell(vr, col).value),
                        "p": num(ws.cell(vr, col + 1).value),
                        "hc": num(ws.cell(vr, col + 2).value),
                        "g": num(ws.cell(vr, col + 3).value),
                    }
                    break
                cant, nota = cantidad(ws.cell(vr, col).value or "")
                ing = {"alimento": clean(h), "cantidad": cant}
                if nota:
                    ing["nota"] = nota
                ingredientes.append(ing)
                col += 1
            rid = slug(nombre_receta(tipo, base))
            r = recetas[tipo].setdefault(rid, {
                "id": rid,
                "nombre": nombre_receta(tipo, base),
                "opcion": clean(base),
                "personas": {},
            })
            pdatos = {}
            if variante:
                plantilla = VARIANTE_NOMBRE.get((tipo, norm(base)))
                pdatos["nombre"] = plantilla.format(v=variante.lower()) if plantilla else f"{r['nombre']} ({variante.lower()})"
            pdatos["ingredientes"] = ingredientes
            pdatos["total"] = total
            r["personas"][pid] = pdatos


def leer_bloques(ws, recetas):
    """COMIDAS Y CENAS: bloques con cabecera 'Alimentos | Cantidades | kcal | ...'."""
    cols_persona = personas_por_columna(ws, 2)
    secciones = []  # (col, 'comida'|'cena')
    for c in range(1, ws.max_column + 1):
        v = ws.cell(4, c).value
        if isinstance(v, str) and v.strip():
            secciones.append((c, "cena" if norm(v).startswith("cena") else "comida"))
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and norm(c.value) == "alimentos"):
                continue
            r0, col = c.row, c.column
            titulo = merged_value(ws, r0, col - 1)
            if not titulo:
                continue
            pid = persona_en(cols_persona, col - 1)
            tipo = persona_en(secciones, col - 1)
            # notas en la fila de encima (p. ej. "Opción2" / descripción)
            notas = []
            for k in range(0, 6):
                v = ws.cell(r0 - 1, col + k).value
                if isinstance(v, str) and v.strip():
                    notas.append(clean(v))
            ingredientes = []
            rr = r0 + 1
            while rr <= ws.max_row:
                nombre = ws.cell(rr, col).value
                if nombre is None:
                    rr += 1
                    continue
                if norm(nombre) == "total":
                    break
                cant, nota = cantidad(ws.cell(rr, col + 1).value or "")
                ing = {
                    "alimento": clean(nombre),
                    "cantidad": cant,
                    "kcal": num(ws.cell(rr, col + 2).value),
                    "p": num(ws.cell(rr, col + 3).value),
                    "hc": num(ws.cell(rr, col + 4).value),
                    "g": num(ws.cell(rr, col + 5).value),
                }
                if nota:
                    ing["nota"] = nota
                ingredientes.append(ing)
                rr += 1
            nombre_app = nombre_receta(tipo, titulo)
            rid = slug(nombre_app)
            r = recetas[tipo].setdefault(rid, {"id": rid, "nombre": nombre_app, "personas": {}})
            notas = [re.sub(r"(?i)\b(opci[oó]n)\s*(\d)", r"\1 \2", n) for n in notas]
            pdatos = {"nota": " · ".join(notas)} if notas else {}
            pdatos["ingredientes"] = ingredientes
            r["personas"][pid] = pdatos


def leer_menu(wb, recetas):
    ws = wb["MENU"]
    dias = []  # (col, nombre)
    for c in range(2, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, str) and v.strip():
            dias.append((c, dia_bonito(v)))
    filas = {}
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and norm(v) in TIPOS:
            filas[norm(v)] = r

    colores = {}
    semanas = []
    for c, dia in dias:
        if dia == "Lunes" or not semanas:
            semanas.append({"nombre": f"Semana {len(semanas) + 1}", "dias": []})
        d = {"dia": dia}
        for tipo, r in filas.items():
            texto = merged_value(ws, r, c)
            if texto is None:
                continue
            d[tipo] = enlazar(tipo, clean(texto), recetas)
            cell = ws.cell(r, c)
            gris = cell.font and cell.font.color is not None and cell.font.color.rgb == "FF999999"
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            if "receta" in d[tipo] and fill and isinstance(fill, str) and not gris and fill not in ("FFF3F3F3", "FFFFFFFF"):
                colores.setdefault(d[tipo]["receta"], "#" + fill[-6:])
        semanas[-1]["dias"].append(d)
    return semanas, colores


def enlazar(tipo, texto, recetas):
    t = norm(ALIAS.get(norm(texto), texto))
    for rid, r in recetas[tipo].items():
        if norm(r["nombre"]) == t:
            return {"receta": rid}
    return {"texto": texto}


def dia_bonito(v):
    return {
        "lunes": "Lunes", "martes": "Martes", "miercoles": "Miércoles",
        "jueves": "Jueves", "viernes": "Viernes", "sabado": "Sábado", "domingo": "Domingo",
    }.get(norm(v), clean(v).capitalize())


# ---------------------------------------------------------------- salida
def volcar(v, nivel=0):
    """JSON legible: objetos "hoja" (sin listas/objetos dentro) en una sola línea."""
    pad, pad2 = "  " * nivel, "  " * (nivel + 1)
    if isinstance(v, float):
        return json.dumps(int(v) if v.is_integer() else round(v, 2))
    if not isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, dict):
        if all(not isinstance(x, (dict, list)) for x in v.values()):
            return "{" + ", ".join(f"{json.dumps(k)}: {volcar(x)}" for k, x in v.items()) + "}"
        items = [f"{pad2}{json.dumps(k, ensure_ascii=False)}: {volcar(x, nivel + 1)}" for k, x in v.items()]
        return "{\n" + ",\n".join(items) + "\n" + pad + "}"
    if not v:
        return "[]"
    if all(not isinstance(x, (dict, list)) for x in v):
        return "[" + ", ".join(volcar(x) for x in v) + "]"
    return "[\n" + ",\n".join(pad2 + volcar(x, nivel + 1) for x in v) + "\n" + pad + "]"


# ---------------------------------------------------------------- main
def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    personas, objetivos = leer_objetivos(wb)
    recetas = {t: {} for t in TIPOS}
    for ws in wb.worksheets:
        n = norm(ws.title)
        if n.startswith("desayuno"):
            leer_opciones_en_fila(ws, "desayuno", recetas)
        elif n.startswith("almuerzo"):
            leer_opciones_en_fila(ws, "almuerzo", recetas)
        elif n.startswith("merienda"):
            leer_opciones_en_fila(ws, "merienda", recetas)
            leer_bloques_si_hay(ws, recetas)
        elif n.startswith("comidas"):
            leer_bloques(ws, recetas)
    semanas, colores = leer_menu(wb, recetas)
    for tipo in TIPOS:
        for rid, r in list(recetas[tipo].items()):
            orden = {"id": r["id"], "nombre": r["nombre"]}
            if r.get("opcion"):
                orden["opcion"] = r["opcion"]
            if rid in colores:
                orden["color"] = colores[rid]  # color del plato en la pestaña MENU
            orden["personas"] = r["personas"]
            recetas[tipo][rid] = orden
    data = {
        "fuente": path.rsplit("/", 1)[-1],
        "actualizado": datetime.date.today().isoformat(),
        "personas": personas,
        "objetivos": objetivos,
        "recetas": {t: list(recetas[t].values()) for t in TIPOS},
        "menu": {"semanas": semanas},
    }
    out = volcar(data)
    print("// Datos de la dieta. Generado desde el Excel con scripts/excel_a_datos.py")
    print("// (se puede editar a mano: cantidades, kcal y macros por ingrediente).")
    print("window.DIETA = " + out + ";")


def leer_bloques_si_hay(ws, recetas):
    """Meriendas en formato bloque (como comidas y cenas), si algún día se añaden así."""
    hay = any(isinstance(c.value, str) and norm(c.value) == "alimentos" for row in ws.iter_rows() for c in row)
    if not hay:
        return
    cols_persona = personas_por_columna(ws, 2)
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and norm(c.value) == "alimentos"):
                continue
            r0, col = c.row, c.column
            titulo = merged_value(ws, r0, col - 1) or f"Merienda {r0}"
            pid = persona_en(cols_persona, col - 1)
            ingredientes, rr = [], r0 + 1
            while rr <= ws.max_row:
                nombre = ws.cell(rr, col).value
                if nombre is None:
                    rr += 1
                    if rr - r0 > 40:
                        break
                    continue
                if norm(nombre) == "total":
                    break
                cant, nota = cantidad(ws.cell(rr, col + 1).value or "")
                ingredientes.append({
                    "alimento": clean(nombre), "cantidad": cant,
                    "kcal": num(ws.cell(rr, col + 2).value), "p": num(ws.cell(rr, col + 3).value),
                    "hc": num(ws.cell(rr, col + 4).value), "g": num(ws.cell(rr, col + 5).value),
                })
                rr += 1
            nombre_app = nombre_receta("merienda", titulo)
            rid = slug(nombre_app)
            r = recetas["merienda"].setdefault(rid, {"id": rid, "nombre": nombre_app, "personas": {}})
            r["personas"][pid] = {"ingredientes": ingredientes}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
